"""Revise research-plan wording in the latest hand-edited CHASI deck.

Only selected text runs and speaker notes change. The source remains untouched;
all drawings, slide layouts, formulas and the topology are retained.
"""
from __future__ import annotations

from hashlib import sha256
from pathlib import Path

from pptx import Presentation


ROOT = Path(__file__).parent
SOURCE = ROOT / 'CHASI_研究汇报_新版方案重做_20261020.pptx'
OUTPUT = ROOT / 'CHASI_研究汇报_研究计划表达_20261020.pptx'
SOURCE_SHA256 = '82fe62a825ef4d794509076b2d1b28b6916332568d2508cb7e33188e1ec6d930'

# slide, shape: (expected original, revised slide text)
REVISIONS = {
    (7, 10): ('尚未接通', '拟研究接口'),
    (7, 12): (
        '只有经过接入与统一评价，才能把“慢层候选 → 快层求解 → 反馈”称为当前联合链。',
        '拟研究慢层候选编码与统一评价接口，再由快层求解；用边缘体验、功率和计算量检验。',
    ),
    (11, 23): (
        '拟研究：由统计误差校准矩阵信道的不确定集合；最坏情况内层未接入。',
        '拟研究：用条件协方差 C 校准链路尺度与方向性椭球，再构造最坏情形内层。',
    ),
    (11, 24): (
        '当前 C 为接收侧 4×4 协方差；新版椭球的坐标域、尺度与置信水平还需定义。',
        '拟按名义、统一和非均匀误差做同条件对照，分析边缘体验、空口功率与求解残差。',
    ),
    (13, 4): (
        '原文用 SRS SINR、距离、速度和 NLOS 构造链路误差方差；函数 f 与量纲尚未标定。',
        '拟用 SRS SINR、距离、速度与 NLOS 拟合链路误差尺度，并校准函数 f 和量纲。',
    ),
    (13, 14): (
        '当前代码用 H∈ℂ⁴×⁶⁴、Cᵣₓ∈ℂ⁴×⁴；扩展为 256 维矩阵误差集合需推导，尚未实现。',
        '基于现有 H∈ℂ⁴×⁶⁴、Cᵣₓ∈ℂ⁴×⁴，研究矩阵信道误差的结构化表征与非均匀椭球。',
    ),
    (14, 16): (
        '任务目标：边缘体验、体验损失 <10%、空口发射功率降 >30%；参照与分母须按原始材料核定，当前无新版效果证据。',
        '拟以边缘体验、<10% 体验损失和 >30% 空口功率目标为参照，开展同条件误差对照。',
    ),
    (15, 4): (
        '上式是原文最坏情形设想；当前统计 C 尚未转化成这个内层问题。',
        '拟以当前统计 C 为基础，研究误差集合校准与最坏情形内层，并与统计平均设计比较。',
    ),
    (15, 10): ('拟加链路组稀疏', '研究链路组稀疏'),
    (15, 12): ('消息门', '拟议消息门'),
    (15, 13): ('m 与缓存比较', '消息变化与缓存比较'),
    (15, 44): (
        '原文尚未定义：本地耦合拆分、发送对象与缓存、触发后共识/对偶更新及收敛条件。',
        '拟研究 TRP 本地变量划分、共识交换与消息缓存；以计算量、消息量和收敛残差检验。',
    ),
}

NOTE_ADDITIONS = {
    7: '研究表达补充：已有基础是两条独立链——历史 GA/PSO 候选生成与评价，以及当前固定 b、χ、r、ν 后的矩阵 WMMSE–ADMM 快层。切入点是把慢层候选映射到正式场景可行的关联/调度配置，定义统一评价输入输出。拟研究候选编码、无效候选修复、快层返回的边缘体验/功率/求解成本以及下一轮候选筛选。验证要在相同场景、测量质量、预算与运行时限下比较，不把虚线解释成已接通调用。',
    11: '研究表达补充：已有基础是每链路接收侧 4×4 条件协方差 C 进入 Ryy 与预编码二次项。切入点是确定从这个统计量到完整 4×64 矩阵误差集合的坐标域、尺度和方向结构；协方差并非自动等于确定性椭球。拟研究置信水平、矩阵向量化、发射侧相关性假设及最坏情形内层的数学形式。验证时固定拓扑、功率预算和优化停止准则，分别比较名义、统一误差与非均匀误差对边缘体验、空口功率及收敛残差的影响；现有对拍不等于这一对照已经完成。',
    13: '研究表达补充：这一页的两层差异是研究入口，不是现成效果。先从现有 Hhat 和 C_rx 的统计估计出发，研究链路间误差尺度由 SRS SINR、CSI 年龄、距离等可观测量如何校准；再研究一条链路内在明确天线或波束域坐标基上的方向结构。当前 Q=Φ/c 是原文经验关系，尚需量纲、分布和置信半径；从单接收 64 维误差向量扩展到 4×64 信道，应明确相关性与向量化次序。验证以观测残差和覆盖概率检查校准，再做同条件资源优化对照；不能预设更大误差会自动展宽波束。',
    14: '研究表达补充：已有基础是 w 已含发射幅度、当前跨资源组共享每 TRP 总功率。研究切入点是明确链路组稀疏和整 TRP 激活两个不同决策层，并给出松弛、阈值判决及重优化后的速率/服务可行性检查。拟研究公平效用、空口功率和稀疏代价的权衡；不能将链路归零说成整台设备休眠，也不能以设备耗电替代空口辐射功率。验证必须以原始难题的边缘体验、平均体验损失和空口功率口径为准，并锁定“80%提升”的共同基线与分母。<10% 和 >30% 是目标，不是新版方案已取得的结果。',
    15: '研究表达补充：右侧状态栏保留当前事实。已有基础是集中式矩阵 WMMSE、统计误差项和 ADMM 可行功率副本。拟研究的事件触发、稀疏与本地计算用相应节点标为“拟议”，不属于当前 ADMM 执行路径。针对 TRP 本地化，需研究跨 TRP 相干项与干扰二次项的变量划分、共识副本、交换消息和缓存；针对触发规则，需区分本轮原始残差、对偶残差以及相对上次发送值的变化，说明静默时继续计算什么。验证比较集中式基线与拟议方法的边缘体验、功率可行性、内外层停止、运行时间和传递量；不预设收敛保证或通信降低比例。',
}


def main() -> None:
    assert sha256(SOURCE.read_bytes()).hexdigest() == SOURCE_SHA256, 'Input changed: re-inspect manual edits.'
    deck = Presentation(SOURCE)
    assert len(deck.slides) == 15
    original_shapes = {
        i: slide.shapes._spTree.xml
        for i, slide in enumerate(deck.slides)
        if i not in {6, 10, 12, 13, 14}
    }
    for (slide_num, shape_num), (old, new) in REVISIONS.items():
        shape = deck.slides[slide_num - 1].shapes[shape_num]
        assert shape.text == old, f'Slide {slide_num}, shape {shape_num} changed.'
        paragraphs = shape.text_frame.paragraphs
        assert len(paragraphs) == 1 and len(paragraphs[0].runs) == 1
        paragraphs[0].runs[0].text = new  # Retain the user's font, size, color and geometry.
    for slide_num, extra in NOTE_ADDITIONS.items():
        notes = deck.slides[slide_num - 1].notes_slide.notes_text_frame
        notes.text = notes.text.rstrip() + '\n\n' + extra
    assert original_shapes == {
        i: deck.slides[i].shapes._spTree.xml for i in original_shapes
    }
    deck.save(OUTPUT)
    print(OUTPUT)
    print('sha256', sha256(OUTPUT.read_bytes()).hexdigest())


if __name__ == '__main__':
    main()
