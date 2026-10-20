"""Build a 15-slide, editable research deck from the supplied PPTX template.

The first two pages are actual template slides. Remaining pages use the
template's regular body layout. Scientific equations are transparent PNGs with
SVG and TeX sources in revision_equations/. Text, tables, arrows, and diagrams
are native editable PowerPoint objects.
"""
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE

HERE=Path(__file__).parent
TEMPLATE=Path(r'C:\Users\sypxx\Desktop\AAAa研一\AAAa研二\开题报告-示范-有调研.pptx')
OUT=HERE/'CHASI_研究汇报_修订版_20261020.pptx'
EQ=HERE/'revision_equations'
TOPO=HERE/'topology'
TOPO_SLIDE=HERE/'revision_topology'
TOPO_SLIDE.mkdir(exist_ok=True)
for source,top,bottom in [('topology_full.png',65,1510),('topology_zoom.png',65,1435)]:
    with Image.open(TOPO/source) as raw:
        # Remove the original plot heading and export/seed caption from the
        # projected page; keep the original image untouched for traceability.
        raw.crop((0,top,raw.width,min(bottom,raw.height))).save(TOPO_SLIDE/source)
BLUE=RGBColor(9,43,86)
RED=RGBColor(177,42,48)
BLACK=RGBColor(0,0,0)
LIGHT=RGBColor(217,230,241)
WHITE=RGBColor(255,255,255)
FONT='Microsoft YaHei'
prs=Presentation(TEMPLATE)

def replace_text(shape, value):
    tf=shape.text_frame
    runs=[r for p in tf.paragraphs for r in p.runs]
    if runs:
        runs[0].text=value
        for r in runs[1:]: r.text=''
    else:
        shape.text=value

# Preserve the template's actual cover and contents slides; discard only the
# remaining example slides. Their masters, layout graphics, and typography stay.
for sid in list(prs.slides._sldIdLst)[2:]:
    prs.part.drop_rel(sid.rId); prs.slides._sldIdLst.remove(sid)
replace_text(prs.slides[0].shapes[0], '非理想测量下的用户中心多TRP协同优化研究')
replace_text(prs.slides[0].shapes[1], '研究汇报')
cover_note=('研究对象：茶思任务中的非理想SRS/CSI与用户中心多TRP协同。'
            '本材料对应已提交代码 db9d7e3；当前另一个会话的工作区更改不包含在本材料。'
            '研究结果尚不能等同正式验收。')
prs.slides[0].notes_slide.notes_text_frame.text=cover_note
for idx,value in [(1,'研究问题与网络'),(2,'系统模型与求解'),
                  (4,'误差机制与验证'),(3,'新版研究方案')]:
    replace_text(prs.slides[1].shapes[idx].shapes[1], value)
prs.slides[1].notes_slide.notes_text_frame.text=(
    '本页给出四段讲述顺序。第一部分界定困难、目标与物理网络；第二部分沿通信变量讲建模与当前求解；'
    '第三部分说明非均匀误差和现有证据；第四部分单列202610版系统方案。'
    '每段均区分原始任务、当前实现、已验证结果与研究设想。')

def text(s,x,y,w,h,value,size=18,bold=False,align=PP_ALIGN.LEFT,color=BLACK):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.015)
    tf.margin_top=tf.margin_bottom=Inches(.015)
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for n,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if n==0 else tf.add_paragraph()
        p.text=line; p.alignment=align; p.space_before=Pt(0); p.space_after=Pt(3)
        for r in p.runs:
            r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return sh

def rule(s,x1,y1,x2,y2,color=LIGHT,width=1.2,dash=False):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=color; sh.line.width=Pt(width)
    if dash: sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
    return sh

def bar(s,x,y,w=0.075,h=.45,color=RED):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=color; sh.line.fill.background(); return sh

def head(s,x,y,w,label):
    bar(s,x,y+.05,h=.35)
    text(s,x+.16,y,w-.16,.45,label,19,True)

def block(s,x,y,w,h,label,body,body_size=17):
    head(s,x,y,w,label)
    text(s,x+.16,y+.51,w-.25,h-.54,body,body_size)

def arrow(s,x1,y1,x2,y2,color=BLUE):
    if abs(y2-y1)<.05:
        sh=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x1),Inches(y1-.10),
                              Inches(max(.22,x2-x1)),Inches(.20))
    else:
        sh=s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x1-.10),Inches(y1),
                              Inches(.20),Inches(max(.22,y2-y1)))
    sh.fill.solid(); sh.fill.fore_color.rgb=color; sh.line.fill.background()
    return sh

def eq(s,name,x,y,w,maxh=.7):
    p=EQ/(name+'.png')
    with Image.open(p) as im: iw,ih=im.size
    h=min(maxh,w*ih/iw); tw=h*iw/ih
    return s.shapes.add_picture(str(p),Inches(x),Inches(y),width=Inches(tw),height=Inches(h))

def new(title,notes):
    s=prs.slides.add_slide(prs.slide_layouts[33])
    if s.shapes.title:
        s.shapes.title.text=title
        s.shapes.title.text_frame.word_wrap=False
        for p in s.shapes.title.text_frame.paragraphs:
            for r in p.runs:
                r.font.name=FONT; r.font.size=Pt(25); r.font.bold=True; r.font.color.rgb=BLUE
    s.notes_slide.notes_text_frame.text=notes
    return s

# 3. Research question and success criteria.
s=new('研究问题与目标',
 '讲解：原始任务的困难是非理想SRS/CSI降低链路辨识，用户中心协同容易让边缘用户受损，同时资源、功率和计算规模受约束。'
 '目标是按原协议核验边缘体验、功率与资源占用，不能直接把SE当体验。'
 '当前方法：固定规则配置资源，WMMSE–ADMM集中求预编码；GA/PSO/SRACOS为独立历史搜索链。'
 '计划：利用链路测量质量建立非均匀误差模型，再公平对照。依据：难题1.png、典型验收场景.png、用户速率协议定义.png；'
 '源码 chasi_run.m、chasi_generate_scene.m、chasi_matrix_wmmse.m。追问：是否达到原验收目标？目前没有正式证明。')
head(s,.62,1.47,3.9,'困难'); text(s,.8,2.05,3.5,1.55,'SRS/CSI 非理想\n边缘 UE 对协同误差更敏感\n多 TRP 带来功率和规模压力',19)
head(s,4.7,1.47,3.9,'研究目标'); text(s,4.88,2.05,3.55,1.55,'按协议比较边缘体验\n满足功率和资源约束\n在正式规模完成可核验求解',19)
head(s,8.76,1.47,3.8,'评价条件'); text(s,8.94,2.05,3.5,1.55,'七站 × 三扇区；273 RB\n每小区 10–20 UE\n体验口径、参考基线需锁定',19)
rule(s,.65,3.91,12.65,3.91)
head(s,.62,4.15,5.6,'当前方法'); text(s,.8,4.77,5.05,1.25,'规则给定调度与服务集合；集中式 WMMSE–ADMM 优化预编码。',19)
head(s,6.55,4.15,5.85,'拟研究创新'); text(s,6.73,4.77,5.5,1.25,'SRS 质量 → 链路差异化误差统计 → 鲁棒优化与公平对照。',19)
text(s,.77,6.25,11.7,.48,'目前已有统计误差机制；新方案的最坏情况、稀疏与事件触发尚待接入和验证。',18)

# 4. MATLAB topology and logical coordination.
s=new('网络拓扑与协同架构',
 '左图是 MATLAB 已导出的真实场景坐标，不是示意站址。中心(450,450)m，外围6站半径300m，21个TRP每站三扇区共址。'
 'UE位置来自当前面积均匀采样，25–150m环扇区为项目假设；扇区图形并非硬覆盖。'
 '右侧为单UE服务和逻辑CPU示意：UE的SRS生成观测，集中仿真得到估计信道和协方差，CPU单进程求w，TRP按块下行。'
 '图像来自 topology_full.png 与 topology_zoom.png；源码 chasi_geometry_from_samples.m、chasi_generate_scene.m、chasi_run.m。'
 '追问：是否TRP分布式ADMM？否。CPU是逻辑节点，图标不代表部署坐标。')
s.shapes.add_picture(str(TOPO_SLIDE/'topology_full.png'), Inches(.45),Inches(1.30),width=Inches(7.25))
head(s,8.04,1.40,4.7,'局部协同与计算');
s.shapes.add_picture(str(TOPO_SLIDE/'topology_zoom.png'), Inches(8.10),Inches(1.93),width=Inches(4.37))
text(s,8.16,5.34,4.34,.54,'UE 7 由 T01、T02 协同服务。',17)
text(s,8.16,5.96,4.34,.98,'SRS → TRP 观测 → CPU 集中求 w\n业务与 w 块 → TRP 下行',16)

# 5. Complete communication-variable chain.
s=new('系统框架',
 '按输入→处理→输出讲六步。几何产生真信道；SRS/CSI产生估计信道及条件误差协方差；资源调度b、服务掩码χ、rank r、RB权重ν由规则给定；'
 '集中WMMSE–ADMM输出含幅度的预编码w；最后以真信道计算用户速率、体验与功率。'
 'Htrue用于合成观测和最终评价，求解接收器依据Hhat/C。依据 chasi_run.m 的场景、优化、评价调用。'
 '追问：优化哪些资源变量？正式入口目前主要优化w。')
stages=[('01 几何','站址、扇区、UE位置','位置与角度'),('02 传播','路径损耗与小尺度衰落','Htrue：真信道'),
        ('03 测量','SRS/CSI 观测与条件更新','Hhat、C：估计和误差'),('04 资源','固定调度、服务、rank','b、χ、r、ν'),
        ('05 求解','WMMSE 外层 / ADMM 内层','w：预编码'),('06 评价','真信道速率和功率统计','体验、功率')]
for i,(h,b,o) in enumerate(stages[:3]+list(reversed(stages[3:]))):
    x=.56+(i%3)*4.29; y=1.61+(i//3)*2.52
    head(s,x,y,3.85,h)
    text(s,x+.18,y+.58,3.62,.74,b,17)
    text(s,x+.18,y+1.43,3.62,.63,'输出：'+o,17,True)
    if i<2: arrow(s,x+3.95,y+1.06,x+4.19,y+1.06)
    elif 3<=i<5:
        sh=s.shapes.add_shape(MSO_SHAPE.LEFT_ARROW,Inches(x+3.94),Inches(y+.96),Inches(.24),Inches(.20))
        sh.fill.solid(); sh.fill.fore_color.rgb=BLUE; sh.line.fill.background()
arrow(s,11.97,3.52,11.97,4.06)
text(s,.75,6.71,11.7,.33,'从观测到求解使用非理想 CSI；性能评价回到真信道。',18)

# 6. Channel and measurement.
s=new('信道与测量模型',
 '几何生成 Htrue∈C^{4×64}；当前是接收相关的几何信道代理，未实现完整双极化物理阵列。'
 'C0=βRrx，CSI-RS与SRS观测在条件高斯假设下精度相加，得到 Cklg∈C^{4×4}、Hhat∈C^{4×64}。'
 '每个发射列共享接收侧4×4条件协方差；未测组省略SRS项。CSI年龄通过噪声膨胀进入。'
 '单位：坐标m，信道为归一化复增益，噪声和协方差相应归一化。依据 CHASI_UNIFIED_TRACE C01–C12，chasi_generate_scene.m。')
head(s,.58,1.48,5.5,'几何与真信道'); eq(s,'channel',.84,2.07,5.42,.7)
text(s,.84,2.92,5.2,1.33,'β：链路大尺度增益；Lrx：接收相关因子；Ξ：小尺度随机矩阵。',18)
head(s,6.47,1.48,6.2,'观测与条件误差'); eq(s,'posterior',6.66,2.07,5.95,.68)
eq(s,'estimate',6.66,3.02,5.95,.66)
text(s,6.66,3.90,5.73,1.45,'C：4×4 条件误差协方差；Hhat：4×64 估计信道。nc、ns 为 CSI-RS 与 SRS 的有效噪声方差。',18)
rule(s,.65,5.66,12.6,5.66)
text(s,.79,5.90,11.65,.78,'测量质量和 CSI 年龄使各链路的 C 不同；当前尚未将 C 解释为已校准的确定性误差椭球。',18)

# 7. Variables, objectives, algorithm division.
s=new('资源变量与优化目标',
 '当前正式链的b、χ、r、ν由规则设定，w为64×r复数预编码且含发射幅度。'
 '功率约束跨资源组：对同一TRP汇总νg加权的所有用户预编码能量，预算Pl默认20W。'
 'PF分支的实际目标为对已调度用户求和log(1+Tk/T0)，其中Tk按资源权重汇总估计SE，T0为PF floor。其他分支还有sum-rate和edge-weighted。'
 'WMMSE 使用估计CSI与误差统计处理物理层目标；最终体验协议是独立评价。'
 '历史候选x为九维归一化编码，分别映射到betaPF、numConnections、scheduleThreshold、rhoLink、rhoPower、maxRank、rankThreshold、repairPower、maxRepairLinks；'
 '这些是旧链参数，不能当作正式场景正在优化的物理资源。GA以交叉、变异和选择产生候选；PSO按位置、速度更新产生候选；'
 'SRACOS承担选择/引导。历史评价器原先以调度流SINR代理为主，存在旧场景条件适配入口；正式chasi_run无搜索调用。'
 '依据 chasi_generate_scene.m、chasi_matrix_wmmse.m、历史搜索源码及 EVIDENCE_INDEX。')
head(s,.57,1.47,5.75,'固定资源规则');
text(s,.78,2.02,5.45,2.05,'b：资源组内 UE 调度；χ：服务 TRP 掩码；r：每 UE 流数，当前为 2；ν：组内发射 RB 比例。',19)
eq(s,'objective_pf',.78,4.06,5.30,.48)
text(s,.80,4.55,5.20,.38,'Tₖ：加权估计速率；T₀：PF 平滑常数。',15.5)
head(s,6.54,1.47,5.98,'当前求解变量');
text(s,6.76,2.02,5.55,1.25,'w：TRP 为 UE 使用的 64×r 复预编码；非服务链路为零。',19)
eq(s,'power',6.76,3.52,5.45,.65)
text(s,6.76,4.26,5.4,.57,'Pₗ：TRP l 的跨组功率预算，单位 W。',17)
rule(s,.62,5.16,12.56,5.16)
head(s,.57,5.35,5.8,'慢层搜索'); text(s,.79,5.91,5.4,.83,'旧 GA/PSO/SRACOS 评价历史候选；未接入正式入口。',17)
head(s,6.55,5.35,5.8,'快层求解'); text(s,6.77,5.91,5.4,.83,'固定 b、χ、r、ν 与 CSI，集中迭代 w。',17)

# 8. Downlink and rate.
s=new('接收模型与速率',
 'w_ljg是TRP l对UE j的预编码块；F_kjg是到UE k的跨TRP相干叠加，有用项为j=k，其余为干扰。'
 'Ryy为4×4接收协方差，I为归一化噪声，p_lg C_klg体现统计误差对接收扰动。Rint=Ryy−FkkFkk^H为4×4干扰加噪声协方差。'
 'SE采用logdet；真信道评价用Htrue重新形成有效信道、干扰与速率。业务体验按原协议另行聚合，不能将SE直接称体验。'
 '当前只有二阶等效计算，没有生成完整OFDM波形。依据 CHASI_UNIFIED_TRACE C15–C20、C32，chasi_matrix_wmmse.m、chasi_evaluate.m。')
head(s,.57,1.47,5.9,'多 TRP 相干叠加'); eq(s,'effective',.77,2.03,5.83,.67)
text(s,.77,2.88,5.47,1.25,'F：4×r 有效信道。先把各服务 TRP 的复数贡献相加，再形成接收功率。',18)
head(s,6.54,1.47,6.05,'接收协方差'); eq(s,'covariance',6.76,2.03,5.79,.68)
text(s,6.76,2.88,5.52,1.26,'Ryy：4×4；含有用信号、所有用户干扰、归一化噪声与误差统计项。',18)
rule(s,.66,4.50,12.56,4.50)
head(s,.57,4.76,6.0,'速率计算'); eq(s,'rate',.78,5.32,5.82,.68)
text(s,6.82,5.33,5.47,1.10,'Rint：干扰加噪声协方差，4×4；RSE：频谱效率，bit/s/Hz。体验另按业务协议统计。',18)

# 9. WMMSE outer loop.
s=new('WMMSE 更新',
 '外层固定当前w，由Hhat/C计算F与Ryy；更新MMSE接收器A∈C^{4×r}，MSE矩阵E∈C^{r×r}，权重Ω=E^{-1}。'
 'η为用户效用/PF权重，固定A、Ω、η后形成w的二次子问题。外层再更新状态并检查目标/残差/迭代限。'
 '速率与WMMSE等价需要固定信道统计、正定E及相应权重条件；联合变量问题不因此变成全局凸。'
 '依据 CHASI_UNIFIED_TRACE C21–C28、chasi_matrix_wmmse.m。')
head(s,.58,1.48,5.8,'输入与接收器');
text(s,.78,2.00,5.52,.68,'固定当前 w，依据估计信道与 C 形成 F、Ryy。',18)
eq(s,'receiver',.79,2.89,5.66,.65)
text(s,.8,3.59,5.34,.60,'A：4×r 接收矩阵，抑制干扰并估计数据流。',17)
head(s,6.57,1.48,5.9,'MSE 与权重'); eq(s,'mse',6.76,2.10,5.72,.73)
eq(s,'weight',6.76,3.15,3.04,.66)
text(s,9.78,3.25,2.64,.62,'Ω：r×r MSE 权重。',17)
rule(s,.61,4.55,12.58,4.55)
head(s,.57,4.78,5.8,'固定量'); text(s,.78,5.33,5.39,.9,'A、Ω、η 固定后，构造预编码二次项。',18)
head(s,6.56,4.78,5.84,'外层返回'); text(s,6.77,5.33,5.43,.95,'ADMM 更新 w → 重算接收器和速率 → 判断停止。',18)

# 10. ADMM.
s=new('ADMM 求解',
 '固定A、Ω、η后，二次系数Q和线性项B由Hhat、C构造。x为预编码原变量，z为满足跨资源组功率约束的副本，y为缩放对偶。'
 '解(νQ+ρ/2 I)x=νB+ρ/2(z−y)，当前用隐式算子与CG；z联合投影全部资源组以满足每TRP预算；随后更新y。'
 '检查原始与对偶残差以及迭代上限。返回可行z作为预编码。当前在单MATLAB进程集中执行，非TRP分布式。'
 '功率可行、数值对拍通过、ADMM收敛和WMMSE外层停止分别判断。依据 CHASI_UNIFIED_TRACE C29–C31、chasi_matrix_wmmse.m。')
head(s,.58,1.42,5.7,'① 原变量：CG 线性求解'); eq(s,'admm_linear',.76,2.02,11.45,.67)
text(s,.78,2.82,11.55,.55,'Q：二次算子；B：线性项；ρ：增广惩罚参数；ν：资源组发射 RB 权重。',18)
rule(s,.63,3.61,12.56,3.61)
head(s,.58,3.88,5.65,'② 可行副本与对偶'); eq(s,'admm_update',.78,4.47,8.70,.68)
text(s,.78,5.23,11.55,.69,'z 对所有资源组联合投影，满足每个 TRP 的跨组预算；y 累积 x 与 z 的差。',18)
rule(s,.63,6.12,12.56,6.12)
text(s,.79,6.31,11.56,.55,'③ 计算原始/对偶残差并判断停止；目前正式规模存在触及迭代上限的记录。',18)

# 11. Nonuniform mechanism.
s=new('非均匀误差机制',
 '链路SRS SINRγs(dB)、CSI年龄ts(s)、先验β与观测选择矩阵P进入有效噪声ns和后验C。'
 '不同链路C不同，作为统计误差项进入Ryy，再改变A/E/Ω和固定接收器后的二次算子Q中的D项，因此可能改变w。'
 '当前实现的是条件协方差统计机制；尚无校准的确定性椭球、最坏情况内层优化，也没有稳定的创新效果正向验证。'
 '公平对照应保持相同拓扑、业务、预算和初始化，只改变名义/统一/非均匀误差模式。依据 chasi_generate_scene.m、chasi_matrix_wmmse.m、CHASI_UNIFIED_TRACE C10–C12/C19/C27。')
head(s,.57,1.45,5.8,'测量质量 → 链路误差'); eq(s,'noise',.77,2.00,5.72,.62)
text(s,.78,2.72,5.42,1.15,'γs：SRS SINR，dB；ts：CSI 年龄，s；τ：相关时间常数，s；ns：有效噪声方差。',18)
head(s,6.57,1.45,5.83,'误差 → 接收与预编码');
text(s,6.79,2.04,5.38,2.03,'C（4×4）进入 Ryy，改变接收器 A、MSE E 和权重 Ω；同时进入预编码二次项 Q，改变 w。',19)
rule(s,.64,4.53,12.56,4.53)
head(s,.57,4.80,5.8,'当前实现'); text(s,.79,5.36,5.32,1.13,'条件协方差及其统计稳健项；还不是确定性最坏情况优化。',18)
head(s,6.57,4.80,5.8,'待验证对照'); text(s,6.79,5.36,5.36,1.13,'名义、统一误差、非均匀误差在同场景下比较边缘体验和功率。',18)

# 12. Dense but readable 2-column glossary.
s=new('核心变量',
 '这是投影用核心符号表。索引 l:TRP、k/j:UE、g:资源组，Nt=64、Nr=4。'
 'Htrue/Hhat均4×64；C4×4；w64×r；F与A为4×r；Ryy4×4；E与Ω为r×r。'
 'x/z/y与w同维，ρ为ADMM惩罚参数，P为W；RSE为bit/s/Hz，Rexp为按业务协议统计的体验率。'
 '完整解释见 SYMBOLS.md 和 CHASI_UNIFIED_TRACE.tex。')
left=[('l,k,j,g','TRP、目标UE、干扰UE、资源组','索引'),('Nt,Nr,r','发射端口、接收天线、数据流数','64、4、2'),
      ('Htrue, Hhat','真信道、估计信道','4×64'),('C','链路条件误差协方差','4×4'),
      ('b, χ, ν','调度、服务掩码、发射RB权重','0/1、0/1、无量纲'),('w','含幅度的预编码','64×r'),
      ('p, Pₗ','组内TRP功率、TRP总预算','W')]
right=[('F','跨TRP相干有效信道','4×r'),('Ryy, A','接收协方差、接收矩阵','4×4、4×r'),
       ('E, Ω, η','MSE、权重、效用','r×r；r×r；标量'),('x,z,y,ρ','原变量、副本、对偶、罚参','同w；ρ为标量'),
       ('RSE','频谱效率','bit/s/Hz'),('Rexp','业务协议体验率','bit/s 或协议定义')]
def table(s,x,y,w,rows):
    widths=[1.36,w-3.1,1.74]
    for j,label in enumerate(('符号','中文含义','维度 / 单位')):
        text(s,x+sum(widths[:j]),y,widths[j],.42,label,16,True)
    rule(s,x,y+.46,x+w,y+.46,BLUE,1.5)
    for i,row in enumerate(rows):
        yy=y+.58+i*.65
        for j,v in enumerate(row): text(s,x+sum(widths[:j]),yy,widths[j],.55,v,15.5)
        if i<len(rows)-1: rule(s,x,yy+.61,x+w,yy+.61,LIGHT,.65)
table(s,.55,1.52,6.16,left)
table(s,6.96,1.52,5.83,right)

# 13. Verified state.
s=new('当前验证结果',
 '已有旧场景的210/420 UE记录：在原始运行参数下能够产生预编码、完成评价并满足功率预算；另有MATLAB/Python数值对拍通过的记录。'
 '这些只证明对应记录中的实现一致性与功率可行性，不证明正式规模收敛，也不证明原始边缘体验验收指标。'
 '最新几何拓扑坐标与服务关系已实际导出；拓扑正确性与求解器收敛是不同层级。'
 '正式运行记录 converged=false，触及外层/内层迭代上限；非均匀优于名义/统一误差的公平对照尚无充分证据。'
 '依据 EVIDENCE_INDEX.md、验证报告与保存日志，所有数字需按同一场景版本解释。')
head(s,.58,1.53,5.75,'已有证据');
text(s,.78,2.10,5.34,2.25,'实际 MATLAB 七站站址、21 TRP 与 UE 坐标\n已有规模的预编码和真信道评价记录\n功率预算检查与数值对拍记录',19)
head(s,6.57,1.53,5.74,'当前不足');
text(s,6.78,2.10,5.33,2.25,'正式规模停止条件尚未达到\n原始边缘体验目标尚未验收\n非均匀误差的增益尚未公平证明',19)
rule(s,.65,5.10,12.56,5.10)
text(s,.78,5.43,11.43,1.18,'功率可行 ≠ 数值对拍通过 ≠ ADMM 收敛 ≠ 外层停止 ≠ 研究目标达成。',20,True)

# 14. V2, separate.
s=new('新版系统方案',
 '此页和下一页仅介绍202610-系统方案v2.pdf的拟议研究，不表示已进入当前代码。'
 '计划从链路级误差椭球入手，结合公平效用、功率、组稀疏和事件触发通信；希望在真实业务与资源限制下评估边缘体验、功率和消息量。'
 '当前统计C与椭球Σ的对应需要校准；联合WMMSE仍非凸，固定部分变量后子问题凸性仍要证明。'
 '关闭某条TRP–UE链路不等于整个TRP休眠。')
head(s,.58,1.50,5.75,'链路不确定性'); eq(s,'robust',.78,2.08,5.65,.62)
text(s,.78,2.83,5.34,1.28,'Σ：拟议椭球尺度；δ：半径。需由测量统计标定，不能直接把当前 C 当作已校准椭球。',18)
head(s,6.57,1.50,5.83,'拟议目标'); eq(s,'proposal',6.76,2.12,5.58,.72)
text(s,6.78,3.03,5.35,1.45,'兼顾公平效用、功率和链路组稀疏；权重 λp、λs 及效用 U 尚须定标。',18)
rule(s,.64,5.08,12.55,5.08)
text(s,.78,5.37,11.54,1.05,'方案是待验证的研究路径；当前代码仍按条件协方差进行集中式预编码优化。',19)

# 15. V2 details and connection boundaries.
s=new('新方案关键机制',
 '讲三条边界：第一，链路级最坏情况需定义椭球与内层优化、论证适用条件；第二，组稀疏需要处理耦合二次项，不能声称整个子问题一次软阈值闭式解决；第三，TRP本地共识与事件触发需要消息缓存、更新时序和收敛条件。'
 '更大误差不必然产生更宽波束；链路关闭不等于TRP休眠；功率下降或通信节省比例是待验证目标。'
 '当前已实现的只有集中式统计误差WMMSE–ADMM及既有评价链，新版系统方案不可画成已完成闭环。'
 '依据 202610-系统方案v2.pdf 与当前 chasi_matrix_wmmse.m。')
for i,(h,b) in enumerate([
 ('误差椭球与最坏情况','先标定链路误差尺度，再定义内层最坏情况问题；联合问题仍需分析凸性。'),
 ('功率与组稀疏','在公平效用外加入功率和链路稀疏代价；关闭链路与TRP休眠分别判定。'),
 ('本地共识与事件触发','定义TRP本地变量、CPU共识、消息缓存与触发条件；以通信量及收敛实测验证。')]):
    y=1.51+i*1.72
    head(s,.59,y,11.9,f'{i+1:02d}  {h}')
    text(s,.81,y+.62,11.35,.76,b,19)
    if i<2: rule(s,.63,y+1.54,12.57,y+1.54)
text(s,.8,6.75,11.46,.32,'当前边界：最坏情况、TRP 分布式与事件触发均未接入正式运行链。',18)

assert len(prs.slides)==15, len(prs.slides)
prs.save(OUT)
print(OUT)
