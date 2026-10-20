"""Rebuild the 202610 proposal section of the latest hand-edited CHASI deck.

Slides 1, 2 and 4 are retained by shape/image signature. The input PPTX is
read-only; this script produces a new 15-slide PPTX plus editable SVG figures
and TeX/math-rendered formula sources. No experiment or solver is run.
"""
from __future__ import annotations

from hashlib import sha256
from pathlib import Path
import importlib.util
import subprocess

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

ROOT=Path(__file__).parent
SOURCE=ROOT/'CHASI_研究汇报_执行链重绘_20261020.pptx'
OUTPUT=ROOT/'CHASI_研究汇报_新版方案重做_20261020.pptx'
SOURCE_SHA='d1eebc340a77a60fa41c33f36ad4a77036f84859f6fcd392a40486caa241bb70'
ART=ROOT/'scheme_v2_sources'
ART.mkdir(exist_ok=True)

# Reuse the already verified native-PowerPoint diagram primitives without
# executing the previous deck generator or touching its input/output files.
spec=importlib.util.spec_from_file_location('prior_draw',ROOT/'rebuild_reference_figures.py')
draw=importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(draw)
draw.scope['FIG']=ART
Diagram=draw.Diagram
clear_body=draw.clear_body
set_title=draw.title
note=draw.note
arrow=draw.arrow
BLUE,MID,GREEN,ORANGE,RED,PALE,VERY=(getattr(draw,k) for k in
    ('BLUE','MID','GREEN','ORANGE','RED','PALE','VERY'))

FORMULAS={
 'original_ellipsoid':r'\Delta\mathbf{h}_{kn}^{H}\mathbf{Q}_{kn}^{-1}\Delta\mathbf{h}_{kn}\leq 1',
 'calibrated_ellipsoid':r'\mathcal{U}_{kn}(\tau)=\{\Delta\mathbf{h}:\Delta\mathbf{h}^{H}(\tau\mathbf{Q}_{kn})^{-1}\Delta\mathbf{h}\leq1\}',
 'proposal_objective':r'\sum_k U_{\alpha}(R_k(\mathbf{w}))-\lambda_{\mathrm{link}}\sum_{l,k}\|\mathbf{w}_{lk,:}\|_F-\lambda_{\mathrm{on}}\sum_l s_l',
 'minmax_core':r'\max_{\mathbf{w}}\min_{\Delta\mathbf{H}\in\mathcal{U}}\sum_k U_{\alpha}(R_k(\mathbf{w},\widehat{\mathbf{H}}+\Delta\mathbf{H}))',
 'shared_power':r'p_l=\sum_g\nu_g\sum_k\|\mathbf{w}_{lkg}\|_F^2\leq s_lP_l',
 'admm_residuals':r'r^{t+1}=\|\mathbf{x}^{t+1}-\mathbf{z}^{t+1}\|,\quad d^{t+1}=\rho\|\mathbf{z}^{t+1}-\mathbf{z}^t\|',
}
def render_formulas():
    import json
    subprocess.run(['py','-3.9',str(ROOT/'render_scheme_formulas.py'),str(ART),
                    json.dumps(FORMULAS,ensure_ascii=False)],check=True)

def add_formula(slide,name,x,y,w,h):
    p=ART/f'{name}.png'
    with Image.open(p) as im:ratio=im.width/im.height
    ww=min(w,h*ratio);hh=ww/ratio
    slide.shapes.add_picture(str(p),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),
                             width=Inches(ww),height=Inches(hh))

def t(d,x,y,w,h,value,size=17,bold=False,color='#000000',center=False):
    return d.text(x,y,w,h,value,size,bold,color,center)

def revise_bridge(slide):
    # Keep the established four-stage communication drawing, but disambiguate
    # its quadratic operator from the proposal's covariance/ellipsoid Q.
    set_title(slide,'现有统计误差机制')
    for sh in slide.shapes:
        if not sh.has_text_frame:continue
        txt=sh.text
        if txt=='二次项 Q(C) → w → 真信道体验':sh.text='二次算子 𝒬(C) → w → 真信道体验'
        elif txt.startswith('待研究：校准椭球'):
            sh.text='拟研究：由统计误差校准矩阵信道的不确定集合；最坏情况内层未接入。'
        elif txt.startswith('公平对照需固定拓扑'):
            sh.text='当前 C 为接收侧 4×4 协方差；新版椭球的坐标域、尺度与置信水平还需定义。'
    note(slide,'本页回答当前代码如何利用非理想 CSI，并为后面的新版方案划清边界。沿箭头从 SRS 质量/年龄讲到接收侧 4×4 条件协方差 C、Ryy 中的误差项、再到固定接收器后的二次算子 𝒬(C) 与预编码 w。这里的 𝒬 是求解器的线性算子，不是新版方案的 64×64 椭球矩阵 Q。当前 awgn_independent 分支把先验、CSI-RS、SRS 的精度相加；pilot_reuse_gaussian 有单独分支。C 只描述统计平均意义的扰动，不保证任何确定性椭球内的最坏速率。下一页是当前符号表，之后单列新方案。依据 db9d7e3 的 chasi_generate_scene.m:65–110、chasi_matrix_wmmse.m:380–444。老师若问是否已做到鲁棒保底：回答是统计误差惩罚已进入计算，最坏情况约束与公平对照结果尚未建立。')

def slide13(slide):
    clear_body(slide)
    # Preserve the user's topic, while shortening the manually edited title
    # so it fits the template's fixed heading box without crossing the logo.
    set_title(slide,'新版方案：测量可靠性驱动协同')
    d=Diagram(slide,'13_measurement_to_ellipsoid',y=1.37,h=5.70)
    t(d,.13,.05,11.95,.43,'SRS 测量误差 → CSI 偏差 → 协同波束与干扰预测失准 → 边缘体验风险',18.1,True,BLUE)
    d.box(.10,.68,5.72,1.34,VERY,MID)
    t(d,.30,.78,5.28,.36,'第一层：不同 UE–TRP 链路的误差规模',17.8,True,BLUE)
    t(d,.30,1.23,5.22,.67,'原文用 SRS SINR、距离、速度和 NLOS 构造链路误差方差；函数 f 与量纲尚未标定。',15.4)
    d.box(6.34,.68,5.73,1.34,VERY,GREEN)
    t(d,6.53,.78,5.28,.36,'第二层：一条链路内的方向差异',17.8,True,BLUE)
    t(d,6.53,1.23,5.27,.68,'Q 的特征值描述各方向误差；对角形式只在选定坐标基中成立。',15.4)
    arrow(d,5.84,1.36,6.30,1.36,MID)
    d.box(.11,2.32,11.95,1.58,PALE,MID)
    t(d,.34,2.40,4.25,.36,'原文的单接收向量椭球',17.0,True,BLUE)
    add_formula(slide,'original_ellipsoid',2.45,4.11,7.70,.53)
    t(d,.35,3.39,11.28,.39,'Δh∈ℂ⁶⁴；Q∈ℂ⁶⁴×⁶⁴。若 Q 是统计协方差，置信边界还需要尺度 τ；原文固定为 1。',15.5)
    d.box(.12,4.06,11.91,.80,'#fff7eb',ORANGE)
    t(d,.34,4.13,11.46,.65,'当前代码用 H∈ℂ⁴×⁶⁴、Cᵣₓ∈ℂ⁴×⁴；扩展为 256 维矩阵误差集合需推导，尚未实现。',16.1,True)
    d.save()
    note(slide,'本页回答为什么提出新版方案，以及“非均匀”究竟有哪两层。从上方因果链讲：SRS 不准使 CSI 有偏差，协同波束和干扰预测可能失准，边缘 UE 的体验风险上升；规模、功率和前传开销也是动机，不是已解决结果。左侧链路尺度来自 202610 Markdown 第31–43行的经验式 σ²=f(SRS SINR,d,v,NLOS)，f、单位和数据标定未给定；C=1/(1+σ²) 只有 σ² 已无量纲化才有清楚含义。右侧是同一链路内方向差异。公式来自本地 1.png，对应 Markdown 第47–51行：Δh 为一条单接收天线链路的 64 维复误差向量，Q 为 64×64 正定形状矩阵；本地 3.png 对应第219–233行的对角 Q。对角仅在天线坐标或指定特征/波束域有效，一般空间相关性不自动对角。若 Q 还代表误差协方差，应引入置信半径 τ，例如 Δh^H(τQ)^−1Δh≤1，并说明分布与置信水平；这属于解释时的补充，不是原文已标定结果。原文 Q=Φ/C 的经验缩放在 C→1 时不趋零，故不能说高置信链路椭球自动收缩到零；C→0 时无界也需截断。当前代码的 H_true/H_est 是 4×64，接收侧 C_rx 是 4×4，列独立假设下完整 vec(H) 协方差形式为 I_64⊗C_rx；它不等于原文发射侧 64×64 Q。下一页讲 w、功率与稀疏决定什么。追问：大误差是否必然展宽波束？不必然，要由目标、信道几何与约束共同决定。来源：原始 Markdown 及本地 1/3.png、chasi_generate_scene.m。')

def slide14(slide):
    clear_body(slide);set_title(slide,'新版方案的优化变量与约束')
    d=Diagram(slide,'14_wave_power_sparse',y=1.35,h=5.80)
    t(d,.11,.05,11.92,.42,'完整波束 w 同时决定发射方向与幅度；功率 p 由波束范数导出。',18.0,True,BLUE)
    d.box(.10,.59,12.03,.82,PALE,MID)
    add_formula(slide,'proposal_objective',.91,2.01,10.92,.57)
    t(d,.23,1.47,11.72,.33,'公平效用 − 链路组稀疏 − TRP 激活代价；分开两种惩罚系数是解释修正。',15.2)
    # Distinct grouping levels, not a generic collection of buzzwords.
    d.box(.12,1.99,3.72,1.67,VERY,MID)
    t(d,.31,2.10,3.33,.40,'波束与功率',18.0,True,BLUE)
    t(d,.31,2.60,3.31,.90,'wₗₖg∈ℂ⁶⁴×ʳ；pₗ 是所有资源组的加权发射功率。',15.7)
    d.box(4.13,1.99,3.73,1.67,VERY,GREEN)
    t(d,4.31,2.10,3.36,.40,'链路组 vs 整 TRP',18.0,True,BLUE)
    t(d,4.31,2.60,3.33,.94,'‖wₗₖ,:‖ 归零只删一条服务链；sₗ=0 才约束整台 TRP 不发射。',15.5)
    d.box(8.17,1.99,3.85,1.67,VERY,ORANGE)
    t(d,8.35,2.10,3.44,.40,'限制与后处理',18.0,True,BLUE)
    t(d,8.35,2.60,3.42,.94,'边缘最低速率、关联可行性；sₗ 松弛判阈后须重优化并复查。',15.5)
    d.box(.14,3.93,11.85,.78,PALE,MID)
    add_formula(slide,'shared_power',1.05,5.44,10.95,.55)
    t(d,.25,4.81,11.60,.42,'任务目标：边缘体验、体验损失 <10%、空口发射功率降 >30%；参照与分母须按原始材料核定，当前无新版效果证据。',15.0,True)
    d.save()
    note(slide,'本页回答新版方案到底优化什么。先说当前预编码 w 已含幅度和方向，不应再无定义地把独立 p 当决策变量；p_l=Σ_gν_gΣ_k||w_lkg||² 是跨资源组/跨 RB 的空口发射功率，s_lP_l 表示 TRP 关断时该 TRP 所有发射块为零。公式的公平效用 Uα、链路稀疏 λ_link、激活成本 λ_on 是对 Markdown 第69–85行原目标的清晰参数化；原文写 λΣ_nΣ_k(||w_nk||+γs_n)，s_n 成本因 Σ_k 被重复计数，分离系数是解释修正，不冒充原文。w_lk,: 的组跨所有 g；该组归零只删 l→k 服务，未必让 l 对其他 UE 停发，更不会消除 l→k 的所有物理干扰。若整 TRP 休眠，可用 s_l=0 和全组功率约束，但需处理固定设备功耗与空口功率不同的口径。原文二进制 s_l 先放松到 [0,1] 再阈值判决，阈值后要在固定活跃集上重优化并检查速率、服务可行性和功率；原文未给出完整过程。α-fair 目标与硬最低速率并存时可行性不能预设。原始“难题1”要求挑战 50% PRB 条件下边缘体验提升达到理想 CSI 集中协同提升的 80%，这更像相对提升比例，不是绝对速率比；共同基线、分母和统计样本需在正式指标定义中核定。<10% 指相对独立功控基线的平均体验损失，>30% 指空口总发射功率降低，不是设备耗电或总能耗。当前代码尚无新版目标实测结果。下一页解释该目标如何可能进入求解器以及缺口。')

def slide15(slide):
    clear_body(slide);set_title(slide,'新版求解流程与实现边界')
    d=Diagram(slide,'15_solver_event_status',y=1.36,h=5.80)
    # Main algorithm spans the left two thirds; status is a compact side ledger.
    d.box(.10,.08,8.34,1.30,PALE,MID)
    t(d,.31,.17,7.94,.34,'WMMSE 外轮：当前 w → A 接收器 → E、Ω 权重 → 固定量求 w',17.1,True,BLUE)
    add_formula(slide,'minmax_core',.69,2.00,7.30,.45)
    t(d,.31,1.10,7.91,.24,'上式是原文最坏情形设想；当前统计 C 尚未转化成这个内层问题。',13.6)
    arrow(d,4.23,1.40,4.23,1.67,MID)
    d.box(.10,1.69,8.34,2.08,'#fff7eb',ORANGE)
    t(d,.31,1.80,7.98,.34,'ADMM 内轮：固定 A、Ω；每轮更新 x → z → y → 双残差',17.0,True,BLUE)
    blocks=[(.32,'x：波束','拟加链路组稀疏'),(2.33,'消息门','m 与缓存比较'),(4.34,'z：副本','联合功率投影'),(6.35,'y：对偶','按 x−z 更新')]
    for x,h,b in blocks:
        d.box(x,2.38,1.78,.90,VERY,ORANGE)
        t(d,x+.09,2.43,1.60,.29,h,15.3,True,BLUE,True)
        t(d,x+.09,2.78,1.60,.43,b,13.4,False,'#000000',True)
    for x in (2.12,4.13,6.14):arrow(d,x,2.84,x+.18,2.84,ORANGE,True)
    t(d,.31,3.34,8.01,.32,'原始残差 ‖x−z‖；对偶残差 ρ‖z⁺−z‖；触发变化 ‖m−缓存‖：三者不同。',14.4,True)
    arrow(d,.24,3.84,.24,4.17,ORANGE)
    d.box(.10,4.19,8.34,.76,PALE,MID)
    t(d,.31,4.26,7.95,.58,'返回可行 z → 检查外层停止；s 阈值判决后重优化并复查速率与功率。',15.4,True)
    d.link(8.45,4.55,8.59,4.55,color=MID,width=2.2)
    d.link(8.59,4.55,8.59,.73,color=MID,width=2.2)
    arrow(d,8.59,.73,8.46,.73,MID)
    # Implementation ledger, explicitly separated from the proposed algorithm.
    d.box(8.79,.09,3.27,4.86,VERY,MID)
    t(d,8.96,.18,2.93,.36,'当前实现状态',17.4,True,BLUE,True)
    status=[('已实现','矩阵 WMMSE、统计 C、集中式 ADMM',GREEN),
            ('部分证据','同输入对拍与功率可行；正式规模未收敛',MID),
            ('未实现','椭球最坏情况、稀疏/休眠、本地 TRP、触发',ORANGE),
            ('未接通','历史 GA/PSO 搜索 → 新求解器',RED)]
    for i,(label,body,col) in enumerate(status):
        y=.69+i*1.00
        d.box(8.97,y,2.90,.30,col,col,.02)
        t(d,9.06,y-.01,2.71,.28,label,14.1,True,'#ffffff',True)
        t(d,9.01,y+.32,2.83,.61,body,12.9)
    t(d,.23,5.12,11.61,.48,'原文尚未定义：本地耦合拆分、发送对象与缓存、触发后共识/对偶更新及收敛条件。',15.3,True,BLUE)
    d.save()
    note(slide,'本页回答 WMMSE 与 ADMM 怎样联合工作，以及稀疏、休眠、事件触发究竟插在哪。先读左上当前外轮：w 形成有效信道，按 Ryy 更新 A、E 和 Ω，固定它们与用户权重 η 后建立预编码二次子问题；这是分块交替，不表示整个联合问题凸。当前统计误差 C 进入 Ryy 与二次算子；原文最坏情形 max_w min_ΔH 见本地 2.png，对应 Markdown 第53–61行，但未给出与 WMMSE 的等价转化，不能把统计平均模型叫作所有椭球误差的保底。再读内轮：当前代码在一个 MATLAB 进程里用 CG 更新 x、联合跨组功率投影 z、缩放对偶 y+=x−z。原文提议把链路组稀疏加到 x 步；近端 group-shrink 的单独算子有闭式形式，但耦合二次项加组稀疏的完整子问题一般不能一次软阈值求解，固定变量后凸性要逐项判断。原文的未缩放 μ 对偶更新应为 μ+=ρ(x−z)，缩放 y=μ/ρ 时为 y+=x−z，当前代码使用后者。原始残差应是更新后 x^{t+1}−z^{t+1}，对偶残差 ρ(z^{t+1}−z^t)；Markdown 第133行写 x^{t+1}−z^t，不能当本轮原始残差。事件触发若按相对上次发送的消息 m 与缓存 m̄ 比较，还必须定义 m 的内容、未发送时协调侧使用的缓存、哪些消息每轮仍发送，以及对偶与共识如何更新；原文第135–141行“静默并跳过对偶”尚不足以形成收敛的完整算法。TRP 本地并行也需拆开跨 TRP 的相干与干扰耦合，变量副本本身不证明已分布式。最后看右侧状态：当前矩阵 WMMSE、接收侧统计 C 和集中 ADMM 有源码；210/420 UE 同输入对拍、功率可行有记录，但正式规模停止条件未通过；最坏情况、组稀疏、休眠、本地计算、消息缓存和事件触发无实现；历史搜索未接正式入口。下一步不是在汇报中承诺效果，而是补充定义和做同条件对照。老师若问 60–80% 消息减少或 30% 功率下降：回答原文是预期，尚无新版实验支持；空口功率不等于设备总能耗。')

def expand_existing_notes(p):
    supplements={
      0:'讲解从茶思原始难题开始：非理想 SRS/CSI 下，多 TRP 协同需要兼顾边缘体验、跨资源功率与规模。标题中的“优化研究”对应当前已实现的矩阵 WMMSE 与集中式 ADMM，也包含尚未完成的鲁棒扩展；不要把本汇报说成完整验收。下一页给出四段阅读顺序。追问：成果是否已达标？目前只能说同输入数值对拍和功率可行有记录，正式规模收敛及新版效果未证实。',
      1:'先让听众看四个分组：研究问题和网络给出对象，系统模型和求解说明如何从信道产生预编码，误差机制和验证说明现有证据，最后新版方案独立介绍拟研究扩展。第 13–15 页的橙色拟议模块不在当前正式执行链。与第 3 页衔接时强调原始验收指标和研究方法是两个不同层次。',
      2:'本页回答“困难为什么需要这些方法”。按每行从左到右读：不同 UE–TRP 的测量质量导致条件误差不同，当前用链路统计 C 进入接收与预编码；边缘 UE 协同受有用信号和干扰共同影响，用用户中心服务和 WMMSE 更新；273 RB 使各资源组共享 TRP 功率预算，当前集中式 ADMM 做联合投影。右侧列的是应检验的量，不是已经达成的效果。当前尚未实现椭球最坏情况或 TRP 本地分布式。第 4 页转到真实七站拓扑。来源为茶思原始三张图、chasi_generate_scene.m 和 chasi_matrix_wmmse.m。老师若问“80% 是否已达成”，回答正式口径仍须锁定共同基线与分母，目前无达成证据。',
      4:'按左到右读真实执行链：几何与业务生成真信道 Htrue；等效 CSI-RS/SRS 质量形成估计 Hhat 和条件误差 C；关联 b、调度 χ、流数 r 和资源权重 ν 由规则固定；矩阵 WMMSE–ADMM 优化含幅度的 w；同一 w 放到 Htrue 上计算速率、体验代理和空口功率。真信道用于仿真生成与最后评价，非理想分支求解使用 Hhat、C。历史搜索链没有调用该正式求解器。下一页拆开测量主体。老师若问“是否已真实协议实现”，回答这是等效测量和集中仿真。',
      5:'从 UE/链路、TRP/接收侧、CPU/协调侧三条泳道读。图中的 SRS 上行、CSI-RS 等效样本只描述仿真观测抽象，不表示完成物理协议栈。Hhat 为 4×64 估计信道，C 为每发射列的 4×4 接收侧条件协方差；CPU 侧根据它们和固定资源求 w，TRP 按 w 发射，UE 再接收。当前预编码和 ADMM 都在单 MATLAB 进程集中计算，不能根据泳道图声称各 TRP 本地并行。下一页解释资源变量。依据 chasi_generate_scene.m、chasi_run.m。',
      6:'先比较两侧：左边历史候选搜索通过 GA 交叉/变异/选择或 PSO 位置/速度更新提议参数，再用旧评价器计分；右边当前正式链固定关联 b、调度 χ、流数 r 和资源权重 ν，只优化预编码 w。橙色虚线表示研究上希望有慢层候选传给快层求解，但当前无实际调用关系。跨资源组功率预算按一个 TRP 共用 P_l，不可每 RB 重新使用整额。下一页从 w 走向真实信道评价。追问：旧候选是否就是正式优化变量？不是，必须先做编码、目标和评价器映射。',
      7:'从发送 w 到真实信道 Htrue，多个 TRP 的同用户流相干叠加；同资源其他用户流形成干扰，叠加噪声得到 UE 接收信号。接收矩阵和 MSE 在优化阶段使用 Hhat、C；最终速率评价重算于 Htrue，不能偷用 Htrue 优化非理想分支。SE 按资源权重转为 TTI 可传比特，再进入业务交付量和用户体验统计。现有体验是代理实现，原协议的有效样本、头尾剔除和底部 5% 口径仍需严格复核。下一页解释求 w 的嵌套算法。',
      9:'本页放大 ADMM 的跨组功率约束和停止判断。固定接收器 A、MSE 权重 Ω 和用户权重 η 后，x 步用 CG 解二次线性系统；z 步对同一 TRP 在所有资源组的加权波束功率联合投影；缩放对偶 y 按 x−z 更新。原始残差、对偶残差、CG 线性求解误差及 WMMSE 外层停止是不同量。当前返回可行副本 z 可以使功率预算满足，但不能证明内层或外层已收敛。210/420 UE 保存运行虽数值对拍与功率可行，正式规模停止条件未通过。下一页看测量误差进入接收计算。',
      11:'本页作查阅符号而非替代每页解释。l、k、g 是 TRP、UE、资源组索引；Htrue/Hhat 是 4×64 真/估计信道；C 是 4×4 接收侧条件误差；w 为 64×r 的完整预编码块。A、E、Ω 是 WMMSE 中的接收器、MSE 和权重；x、z、y 是 ADMM 原变量、可行副本、缩放对偶。p 是 w 的加权范数，P 是每 TRP 预算；RSE 和 Rexp 分别为物理频谱效率与业务体验统计口径。新方案的 64×64 椭球 Q 不应与现有 4×4 C 或二次算子 𝒬 混淆。下一页据此进入拟研究方案。'
    }
    for idx,extra in supplements.items():
        frame=p.slides[idx].notes_slide.notes_text_frame
        old=frame.text.strip()
        frame.text=(old+'\n\n' if old else '')+extra

def main():
    assert sha256(SOURCE.read_bytes()).hexdigest()==SOURCE_SHA,'Hand-edited input changed; re-inspect it.'
    render_formulas()
    p=Presentation(SOURCE)
    assert len(p.slides)==14
    kept={i:p.slides[i].shapes._spTree.xml for i in (0,1,3)}
    revise_bridge(p.slides[10])
    slide13(p.slides[12]);slide14(p.slides[13])
    final=p.slides.add_slide(p.slides[13].slide_layout)
    slide15(final)
    expand_existing_notes(p)
    assert len(p.slides)==15
    assert kept=={i:p.slides[i].shapes._spTree.xml for i in kept}
    p.save(OUTPUT)
    print(OUTPUT,sha256(OUTPUT.read_bytes()).hexdigest())

if __name__=='__main__':main()
