"""Enhance the user's manually edited 14-slide PPTX in place *in memory*.

The input is never saved. Slides 1–2 and the hand-positioned topology pictures
on slide 4 are retained; the result is written under a new name.
"""
from pathlib import Path
import hashlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from PIL import Image

ROOT = Path(__file__).parent
SOURCE = ROOT / 'CHASI_研究汇报_修订版_20261020.pptx'
OUT = ROOT / 'CHASI_研究汇报_图示增强版_20261020.pptx'
EQ = ROOT / 'revision_equations'
EXTRA_EQ = ROOT / 'enhanced_equations'
EXPECTED = '61ae41073d08d4a796332a82ea09ff7006b2a7d3cd8b6a199f1b42a230e49926'
actual = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
if actual != EXPECTED:
    raise RuntimeError(f'Hand-edited source has changed; inspect before rebuilding: {actual}')

prs = Presentation(SOURCE)
assert len(prs.slides) == 14
BLUE = RGBColor(9, 43, 86)
MID = RGBColor(18, 74, 143)
PALE = RGBColor(224, 236, 246)
PALE2 = RGBColor(244, 249, 253)
RED = RGBColor(177, 42, 48)
ORANGE = RGBColor(232, 121, 41)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
FONT = 'Microsoft YaHei'

def keep_title(slide):
    """Remove generated page content while keeping the template title."""
    # The supplied template's heading is a placeholder whose type is not
    # reported by python-pptx as shapes.title. It is still shape 0.
    title = slide.shapes[0]
    for shape in list(slide.shapes):
        if shape._element is not title._element:
            slide.shapes._spTree.remove(shape._element)
    return slide

def txt(s, x, y, w, h, value, size=18, bold=False, color=BLACK,
        align=PP_ALIGN.LEFT, fill=None, margin=.04):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(.015)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(str(value).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_before = Pt(0); p.space_after = Pt(2)
        for run in p.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    return sh

def rect(s, x, y, w, h, color, rounded=False, outline=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color
    if outline:
        sh.line.color.rgb = outline; sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    return sh

def line(s, x1, y1, x2, y2, color=BLUE, width=1.4, dash=False):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                Inches(x2), Inches(y2))
    sh.line.color.rgb = color; sh.line.width = Pt(width)
    if dash: sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return sh

def arrow(s, x, y, w, h, color=MID):
    sh = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = color; sh.line.fill.background()
    return sh

def heading(s, x, y, w, label, size=20):
    rect(s, x, y+.11, .07, .31, RED)
    txt(s, x+.17, y, w-.17, .53, label, size, True)

def eq(s, name, x, y, w, maxh=.68):
    p = EQ / (name + '.png')
    if not p.exists(): p = EXTRA_EQ / (name + '.png')
    with Image.open(p) as im:
        h = min(maxh, w*im.height/im.width)
        realw = h*im.width/im.height
    s.shapes.add_picture(str(p), Inches(x), Inches(y), width=Inches(realw), height=Inches(h))

def note(s, body):
    s.notes_slide.notes_text_frame.text = body

# Slide 1 and 2 are deliberately untouched. Their actual template layouts and
# the user's current wording survive intact.

# 3: research problem as a pictorial mechanism, not five parallel text boxes.
s = keep_title(prs.slides[2])
heading(s,.56,1.28,5.85,'非理想测量使协同决策变难')
txt(s,.76,1.79,5.40,.63,'同一 UE 与不同 TRP 的 CSI 精度不同',19)
for x,label,col in [(1.06,'TRP 1',MID),(3.31,'TRP 2',ORANGE)]:
    rect(s,x,2.69,1.40,.56,col,True)
    txt(s,x,2.69,1.40,.56,label,17,True,WHITE,PP_ALIGN.CENTER)
rect(s,2.25,4.05,1.23,.58,PALE,True,MID)
txt(s,2.25,4.05,1.23,.58,'边缘 UE',17,True,BLACK,PP_ALIGN.CENTER)
line(s,1.74,3.25,2.64,4.04,MID,2.4)
line(s,4.01,3.25,3.17,4.04,ORANGE,2.4,True)
txt(s,.91,3.44,1.37,.43,'测得较准',17,True,MID)
txt(s,3.69,3.44,1.47,.43,'测得较差',17,True,ORANGE)
txt(s,.78,5.22,5.37,.79,'误差改变相干合成与干扰估计，边缘体验最容易受影响。',18)
txt(s,.79,6.13,5.22,.40,'用户中心协同研究背景 [1,2]',15.9)
line(s,6.43,1.52,6.43,6.47,PALE,1.2)
heading(s,6.73,1.28,5.93,'目标、方法与验证')
for i,(n,h,b) in enumerate([
    ('1','研究目标','按原协议评估边缘体验，并满足跨资源组功率和规模要求'),
    ('2','当前方法','规则确定服务与调度；集中式 WMMSE–ADMM 求预编码'),
    ('3','拟研究创新','SRS 质量驱动链路差异化误差；同条件公平对照')]):
    y=2.00+i*1.45
    rect(s,6.84,y,.48,.48,RED,True)
    txt(s,6.84,y,.48,.48,n,18,True,WHITE,PP_ALIGN.CENTER)
    txt(s,7.49,y,2.1,.42,h,19,True)
    txt(s,7.48,y+.48,5.03,.68,b,16.6)
note(s,'回答：研究困难、目标、现有求解和拟研究创新分别是什么。左侧两条链路为机制示意，不是实际拓扑或测量值。原始材料难题1.png、典型验收场景.png、用户速率协议定义.png要求关注边缘体验、功率和规模。用户中心服务与可扩展性背景见[1][2]。当前资源由规则固定、预编码由集中式 WMMSE–ADMM 求解；最坏情况鲁棒优化未接入。下一页给出真实场景坐标。[1] Demir, Björnson, Sanguinetti, Foundations of User-Centric Cell-Free Massive MIMO, 2021, DOI:10.1561/2000000109. [2] Ammar et al., Distributed Resource Allocation Optimization for User-Centric Cell-Free MIMO Networks, IEEE TWC 2022, DOI:10.1109/TWC.2021.3118303。追问：是否满足验收体验？现有记录尚不足以证明。')

# 4: retain the user's two hand-positioned MATLAB images exactly. Their dots
# are colored by initial home site, not by scheduling or full serving set.
s=prs.slides[3]
manual_pictures=[sh for sh in s.shapes if sh.shape_type == 13]
assert len(manual_pictures)==2
picture_geometry=[(sh.left,sh.top,sh.width,sh.height) for sh in manual_pictures]
for sh in list(s.shapes):
    if sh._element is not s.shapes[0]._element and all(sh._element is not p._element for p in manual_pictures):
        s.shapes._spTree.remove(sh._element)
heading(s,7.80,1.27,4.92,'共址扇区 · 用户中心服务')
txt(s,8.14,5.37,4.62,.55,'① 三个 TRP 共用一个站址，箭头表示扇区主方向。',16.4)
txt(s,8.14,5.93,4.62,.55,'② UE 7 由 T01、T02 服务；CPU 集中计算预编码。',16.4)
txt(s,8.14,6.47,4.60,.37,'彩色小点＝实际 UE；颜色＝初始归属站址。',15.4,True)
assert picture_geometry==[(sh.left,sh.top,sh.width,sh.height) for sh in manual_pictures]
note(s,'回答：七站三扇区、UE 空间分布、服务关系和 CPU 分工。左图是 MATLAB 输出的真实坐标；全部210个 UE 均为独立散点，未重复画点。MATLAB lines(7)颜色按 home 所在站 S1–S7：S1蓝、S2橙、S3黄、S4紫、S5绿、S6青、S7洋红/暗红。颜色不是调度或完整服务集合。右图 UE7 的 home=T01，服务集合 T01/T02；两TRP共址，服务连线可能重合。CPU 为逻辑协调节点，不是实际部署坐标。原材料规定每小区10–20 UE、小区内均匀分布；当前10 UE/小区，总210。chasi_geometry_from_samples.m 以中心(450,450)m、外围6站距中心300m生成；每站三个 TRP 共址；UE 在每扇区25–150m环扇区，以 θ=φ_l+(v−1/2)2π/3 与 r=sqrt(rmin²+u(rmax²−rmin²))采样，因此面积均匀。25–150m边界是项目假设。home 为生成扇区，服务集合为 home 加最近活动 TRP，当前上限2。扇形不是无干扰硬覆盖。下一页沿变量看执行链。追问：紫色点是什么？是初始归属 S4 的真实 UE。')

# 5: template page 3 inspired continuous blue chevrons.
s=keep_title(prs.slides[4])
stages=[('场景与信道','位置、扇区','真信道 Hᵗʳᵘᵉ'),('测量与估计','SRS / CSI','估计 Ĥ、误差 C'),
        ('资源配置','归属与规则','调度 b、服务 χ、rank r、权重 ν'),
        ('预编码求解','WMMSE–ADMM','预编码 w'),('性能评价','真信道计算速率','体验率与功率')]
for i,(h,b,o) in enumerate(stages):
    x=.55+i*2.55
    arrow(s,x,1.55,2.43,.92,BLUE if i<3 else MID)
    txt(s,x+.08,1.71,1.99,.57,h,17.3,True,WHITE,PP_ALIGN.CENTER)
    txt(s,x+.06,2.78,2.27,.62,b,17,True,BLACK,PP_ALIGN.CENTER)
    line(s,x+1.18,2.49,x+1.18,2.72,PALE,1.3)
    txt(s,x+.02,3.55,2.35,1.14,o,16.5,False,BLACK,PP_ALIGN.CENTER)
    if i<4: line(s,x+2.43,3.08,x+2.55,3.08,MID,1.7)
heading(s,.61,5.12,6.12,'变量传递的关键边界')
txt(s,.79,5.70,11.65,.73,'求解使用 Ĥ 与 C；调度和服务集由规则固定；最终用 Hᵗʳᵘᵉ 计算实际速率。',19)
note(s,'回答：当前正式入口如何从场景产生预编码并评价。沿五个蓝色大箭头讲输入→处理→输出。Htrue∈C^{4×64}用于合成观测和真实评价；Hhat、C进入集中式求解；b、χ、r、ν由规则配置，w是当前主优化变量。求解器不调用历史GA/PSO/SRACOS链。源码 chasi_run.m、chasi_generate_scene.m、chasi_matrix_wmmse.m、评价函数。下一页展开场景与测量。')

# 6: exact geometry and channel structure confirmed against MATLAB source.
s=keep_title(prs.slides[5])
heading(s,.57,1.30,5.72,'七站几何与 UE 采样')
eq(s,'geometry',.77,1.92,5.61,.66)
txt(s,.81,2.75,5.43,.74,'中心 (450,450) m；外围 6 站相隔 60°，距中心 300 m。',17.5)
rect(s,.75,3.65,5.40,1.63,PALE2,True,PALE)
txt(s,.90,3.79,5.13,.42,'UE 在所属扇区的环形区域内面积均匀抽样',17,True)
eq(s,'ue_radius',.90,4.31,5.05,.53)
txt(s,.85,5.54,5.26,.72,'半径 25–150 m 与扇区 ±60° 是项目假设。',17)
line(s,6.47,1.52,6.47,6.43,PALE,1.2)
heading(s,6.68,1.30,5.75,'当前真信道：接收相关模型')
eq(s,'channel',6.86,1.93,5.57,.66)
eq(s,'rx_factor',6.86,2.88,5.52,.61)
txt(s,6.88,3.71,5.35,1.21,'β：大尺度增益；Lᵣₓ：4×4 接收相关因子；Ξ：4×64 独立复高斯矩阵。',17.3)
heading(s,6.69,5.18,5.56,'测量输出')
txt(s,6.88,5.75,5.32,.74,'SRS/CSI 观测形成 Ĥ（4×64）与条件误差 C（4×4）。',17.4)
note(s,'回答：第六页公式是否错误、UE怎样采样。chasi_geometry_from_samples.m：七站 s_c=(450,450)m，外围 s_m=s_c+300[cos(mπ/3),sin(mπ/3)]，m=0,…,5。21TRP由repelem(site,3,1)共址生成。home=repelem((1:21),10)；θ=mod(l−1,3)2π/3+(v−1/2)2π/3，v∼U(0,1)；当前 ueSpatialSampling=area_uniform_annular_sector，r=sqrt(rmin²+u(rmax²−rmin²))，rmin25m、rmax150m。源码确实按面积均匀采样，旧linear-with-sqrt分支不是当前配置。验收图仅要求每小区10–20UE且小区内均匀分布，未指定当前环形边界或中心坐标。chasi_generate_scene.m：rxCorrelation=0.35.^abs(i−j)，rxChol=chol(rxCorrelation,lower)，truth=sqrt(gain)*rxChol*complex_normal(4,64)，后者每元素(randn+i randn)/sqrt(2)。因此 Htrue=√β Lrx Ξ、Lrx Lrxᴴ=Rrx、Ξ(a,m)∼CN(0,1)，维度4×64；原版乘法方向和维度正确，没有发现实际公式错误，本版补足缺少的定义。没有发射相关或 CDL 模型。下一页区分固定资源和优化变量。')

# 7: resource variables, objective, constraints, and slow/fast split.
s=keep_title(prs.slides[6])
heading(s,.57,1.34,5.72,'规则固定的资源')
for i,(symbol,meaning) in enumerate([('b','UE 是否调度'),('χ','TRP–UE 服务掩码'),('r','数据流数 / rank'),('ν','资源组占用 RB 数')]):
    y=1.98+i*.81
    rect(s,.78,y,1.03,.55,PALE,True)
    txt(s,.78,y,1.03,.55,symbol,20,True,BLUE,PP_ALIGN.CENTER)
    txt(s,2.00,y,3.93,.55,meaning,18)
line(s,6.48,1.51,6.48,6.5,PALE,1.2)
heading(s,6.70,1.34,5.75,'当前优化的连续变量')
txt(s,6.91,2.02,5.19,.80,'wₗₖg：TRP l 对 UE k 的预编码，64×r，包含发射幅度。',18)
eq(s,'objective_pf',6.89,3.04,5.40,.58)
txt(s,6.91,3.76,5.20,.66,'效用以交付量的 PF 代理计算；不同于最终体验率。',17.1)
eq(s,'power',6.88,4.59,5.45,.63)
txt(s,6.91,5.37,5.18,.68,'每个 TRP 的功率在所有资源组上合计受限。',17.1)
line(s,.63,6.39,12.56,6.39,PALE,1.1)
txt(s,.79,6.53,11.44,.41,'历史 GA / PSO / SRACOS 搜索链仍独立；未调用当前矩阵 WMMSE 评价器。',17.1)
note(s,'回答：哪些是资源规则，哪些是当前优化量。b调度、χ服务集合、r流数、ν资源组RB权重由固定规则生成；当前集中式矩阵WMMSE–ADMM优化w。PF代理与业务协议体验率不是同一个目标。功率约束∑gνg∑j||w_ljg||²≤P_l，跨资源组联合投影。历史搜索的候选编码/评价器与正式矩阵求解链未接通，不能画成实际两时间尺度闭环。依据 chasi_run.m、chasi_matrix_wmmse.m 与公式合集。下一页讲速率链。追问：GA/PSO是否优化正式w？没有当前调用证据。')

# 8: signal flow + equations.
s=keep_title(prs.slides[7])
heading(s,.57,1.30,5.77,'用户中心相干信号')
for i,(lab,sub) in enumerate([('TRP 服务块','Hₖₗg · wₗⱼg'),('跨 TRP 合成','Fₖⱼg'),('UE 接收','有用 + 干扰 + 噪声')]):
    x=.71+i*4.15
    rect(s,x,1.98,3.37,.91,PALE2,True,PALE)
    txt(s,x+.13,2.04,3.08,.36,lab,18,True,BLUE,PP_ALIGN.CENTER)
    txt(s,x+.12,2.43,3.07,.30,sub,16.4,False,BLACK,PP_ALIGN.CENTER)
    if i<2: arrow(s,x+3.52,2.32,.49,.22,MID)
eq(s,'effective',.76,3.33,5.85,.62)
txt(s,.83,4.15,5.36,.71,'Fₖⱼg（4×r）：UE k 所见 UE j 的相干等效信道。',17.3)
line(s,6.45,3.18,6.45,6.54,PALE,1.2)
heading(s,6.68,3.19,5.72,'干扰、速率与评价')
eq(s,'covariance',6.84,3.80,5.39,.58)
txt(s,6.85,4.51,5.32,.70,'Rʸʸ（4×4）＝信号、干扰与误差统计的接收协方差。',16.8)
eq(s,'rate',6.85,5.39,5.39,.56)
txt(s,6.86,6.05,5.32,.55,'SE 用 bit/s/Hz；体验率按业务协议另算。',16.8)
note(s,'回答：预编码怎样变成可评价的速率。TRP l 的 w_ljg 与 H_klg相乘，按用户中心服务掩码跨TRP相干叠加为 F_kjg∈C^{4×r_jg}；目标用户有用项 F_kkg，其他j产生干扰。Ryy含噪声、全部用户信号和统计误差∑l p_lg C_klg，维度4×4。由干扰加噪声协方差计算RSE，之后结合RB与业务交付得到体验统计，不能将SE直接当最终体验率。公式和符号来自当前源码及FORMULA_TRACE。下一页说明WMMSE更新。')

# 9: outer iteration drawn as a cycle with equations.
s=keep_title(prs.slides[8])
heading(s,.57,1.31,5.85,'外层：交替更新接收与预编码')
steps=[('当前 w','含幅度预编码'),('A ← 接收器','4×r'),('E、Ω ← MSE','r×r'),('w ← ADMM','功率可行副本')]
for i,(lab,sub) in enumerate(steps):
    x=.64+i*3.14
    rect(s,x,1.96,2.69,.87,PALE2,True,PALE)
    txt(s,x+.08,2.01,2.52,.35,lab,17.2,True,BLUE,PP_ALIGN.CENTER)
    txt(s,x+.08,2.39,2.52,.30,sub,15.7,False,BLACK,PP_ALIGN.CENTER)
    if i<3: arrow(s,x+2.76,2.27,.32,.18,MID)
line(s,11.94,3.08,11.94,3.33,RED,1.8)
line(s,11.94,3.33,1.98,3.33,RED,1.8)
txt(s,5.04,3.41,3.17,.37,'重算并判断外层停止',16.3,True,RED,PP_ALIGN.CENTER)
eq(s,'receiver',.78,4.08,5.85,.60)
eq(s,'mse',6.68,4.08,5.79,.59)
eq(s,'weight',.78,5.33,3.70,.55)
txt(s,4.71,5.32,7.42,.68,'η：用户效用权重；固定 A、Ω、η 后，构造 w 的二次子问题。',17.3)
txt(s,.80,6.22,11.47,.46,'输入 Ĥ、C 与当前 w；输出更新后的 A、E、Ω，并继续优化 w。',17.2)
txt(s,10.27,5.73,2.19,.36,'理论依据 [4]',15.9)
note(s,'回答：WMMSE外层的输入、固定量、更新量。以当前w和Hhat/C形成F、Ryy；A=(Ryy)^{-1}F_kkg，A∈C^{4×r}；E∈C^{r×r}来自接收误差公式；Ω=E^{-1}∈C^{r×r}。η是由当前PF目标对应的用户效用权重。固定A、Ω、η后构造预编码二次项Q和B，进入ADMM；返回w再重算速率/目标和停止指标。WMMSE理论源于[4] Shi et al., IEEE TSP 2011, DOI:10.1109/TSP.2011.2147784；当前实现按项目目标调整，并非直接复制论文代码。联合问题仍非凸。下一页展开ADMM。')

# 10: ADMM algorithm as a visually distinctive loop.
s=keep_title(prs.slides[9])
heading(s,.56,1.29,5.89,'内层：固定 A、Ω、η，更新 w')
eq(s,'admm_linear',.74,1.89,11.85,.63)
txt(s,.78,2.60,11.48,.53,'x：原变量；Q、B：二次项；ρ：罚参数；νg：资源组 RB 权重。',17.2)
for i,(h,b) in enumerate([('① CG 求 x','隐式二次算子'),('② 投影 z','跨组功率可行'),('③ 对偶 y','累积 x−z'),('④ 残差','判定内层停止')]):
    x=.66+i*3.17
    rect(s,x,3.38,2.77,1.26,PALE2,True,PALE)
    txt(s,x+.10,3.49,2.56,.41,h,17.6,True,BLUE,PP_ALIGN.CENTER)
    txt(s,x+.10,3.97,2.56,.48,b,16.4,False,BLACK,PP_ALIGN.CENTER)
    if i<3: arrow(s,x+2.80,3.86,.31,.18,MID)
eq(s,'admm_update',.77,5.14,8.02,.62)
eq(s,'power',.77,6.06,8.90,.57)
txt(s,9.56,5.25,2.52,1.24,'返回功率可行的 z；与数值对拍、外层停止分别检查。',16.2)
txt(s,9.59,6.61,2.37,.33,'方法依据 [5]',15.9)
note(s,'回答：ADMM实际解什么。固定WMMSE接收器和权重后，形成组g二次项Q_g与B_g，原变量线性系统为(ν_gQ_g+ρ/2 I)x_g=ν_gB_g+ρ/2(z_g−y_g)，当前用隐式算子和CG。z=Π_C(x+y)是满足∀l:∑gν_g∑j||w_ljg||²≤P_l的联合投影，因此不能对每组独立投影；y←y+x−z。检查原始和对偶残差以及迭代上限，返回可行z。代码为单MATLAB进程集中式实现，使用ADMM并不意味着TRP分布式。正式规模运行有触及迭代上限的记录；功率可行、数值对拍、内层残差收敛、外层停止是不同事实。[5] Boyd et al., Distributed Optimization and Statistical Learning via ADMM, 2011, Stanford作者页：https://stanford.edu/~boyd/papers/admm_distr_stats.html。下一页看误差如何改变Q。')

# 11: three-level mechanism and comparison, distinguish statistical C/ellipsoid.
s=keep_title(prs.slides[10])
heading(s,.57,1.30,5.85,'链路测量质量 → 非均匀统计误差')
for i,(h,b) in enumerate([('测量输入','SRS SINR、CSI 年龄、先验 β'),('条件统计','各链路 Cₖₗg：4×4'),('进入求解','Rʸʸ → A、E、Ω；Q → w')]):
    x=.69+i*4.17
    rect(s,x,2.01,3.48,1.10,PALE2,True,PALE)
    txt(s,x+.12,2.08,3.22,.42,h,18.2,True,BLUE,PP_ALIGN.CENTER)
    txt(s,x+.13,2.57,3.21,.42,b,16.3,False,BLACK,PP_ALIGN.CENTER)
    if i<2: arrow(s,x+3.55,2.43,.47,.20,MID)
eq(s,'noise',.78,3.66,6.16,.64)
txt(s,7.12,3.66,5.08,1.05,'γs（dB）、CSI 年龄 t（s）改变有效噪声；后验 C 因链路而异。',17.3)
line(s,.69,4.84,12.52,4.84,PALE,1.1)
heading(s,.57,5.01,5.85,'公平对照与当前边界')
txt(s,.81,5.63,11.44,.75,'同一拓扑、业务、预算与初始化下比较：名义 CSI ｜统一误差｜非均匀误差。',18.2)
txt(s,.81,6.35,11.44,.47,'当前是统计协方差模型；正式规模尚未达停止条件；最坏情况内层未接入。',16.8,True)
note(s,'回答：为何非均匀误差会进入预编码，而非只在报告里改口径。chasi_generate_scene.m依据SRS SINR、CSI年龄、先验β和观测选择结构计算后验条件误差协方差C_klg（4×4）。各链路C不同。chasi_matrix_wmmse.m把∑l p_lg C_klg放入Ryy，故A、E、Ω随之变；固定A/Ω后的预编码二次算子Q亦含误差贡献，因此w可能改变。当前C是条件协方差，不是已校准的确定性误差椭球，当前无最坏情况内层。公平对照需同拓扑/业务/功率预算/初始化，仅切换误差假设；现有记录尚未证明非均匀模式提升。非理想CSI背景见[2] Ammar et al. 2022；椭球鲁棒性的未来理论来源[6] Zheng et al. 2008 DOI:10.1155/2008/609028。下一页汇总核心符号。')

# 12: template page 6 inspired grouped blue/pale table.
s=keep_title(prs.slides[11])
columns=[.72,2.52,7.68,10.94,12.57]
for j,label in enumerate(['类别','标准符号','中文含义','维度 / 单位']):
    rect(s,columns[j],1.48,columns[j+1]-columns[j],.49,BLUE)
    txt(s,columns[j]+.05,1.49,columns[j+1]-columns[j]-.10,.47,label,16.3,True,WHITE,PP_ALIGN.CENTER)
rows=[
('索引与规模','l, k, j, g; Nₜ, Nᵣ, r','TRP、UE、资源组；端口、接收天线、流数','索引；64、4、≤2'),
('信道与测量','Hᵗʳᵘᵉ, Ĥ, C','真信道、估计信道、条件误差协方差','4×64；4×64；4×4'),
('资源变量','b, χ, ν; w, p, Pₗ','调度、服务、RB 权重；预编码、功率、预算','0/1；64×r；W'),
('WMMSE','F, Rʸʸ, A, E, Ω, η','有效信道、接收协方差、接收器、MSE、权重、效用','4×r；4×4；r×r'),
('ADMM','x, z, y, ρ','原变量、可行副本、缩放对偶、罚参数','同 w；标量'),
('评价','Rˢᴱ, Rᵉˣᵖ','频谱效率、按协议统计的体验率','bit/s/Hz；协议口径')]
for i,row in enumerate(rows):
    y=2.04+i*.75
    rect(s,columns[0],y,columns[-1]-columns[0],.69,PALE if i%2==0 else PALE2)
    for j,v in enumerate(row):
        txt(s,columns[j]+.07,y+.04,columns[j+1]-columns[j]-.14,.59,v,
            15.2 if j!=2 else 14.9,j==0,BLACK)
    if i<len(rows)-1: line(s,columns[0],y+.72,columns[-1],y+.72,WHITE,.7)
note(s,'回答：跨图使用的核心符号各表示什么，哪些有物理单位。维度r随用户与资源组变化，表中符号省略下标。Htrue/Hhat∈C^{4×64}，C、Ryy∈C^{4×4}，w∈C^{64×r}，A/F∈C^{4×r}，E/Ω∈C^{r×r}；x/z/y与w同维。ν为RB计数/权重，p和P_l按当前归一化功率制说明；SE=bit/s/Hz，Rexp依业务协议和采样规则定义。完整符号见现有FORMULA_TRACE与SYMBOLS.md。辅助索引、维度及假设可口头解释。下一页只介绍独立新版方案。')

# 13–14: preserve user's decision to keep V2 separate, but use two distinct
# visuals rather than another duplicated left/right text layout.
s=keep_title(prs.slides[12])
heading(s,.57,1.28,5.79,'从测量统计到拟议机制')
for i,(h,b) in enumerate([('链路椭球','由测量统计标定不确定性集合'),
                          ('公平效用','结合边缘体验与功率成本'),
                          ('组稀疏','减少弱效链路的服务开销'),
                          ('事件触发','在保证更新条件下减少通信')]):
    y=1.96+i*1.13
    rect(s,.81,y,.51,.51,RED,True)
    txt(s,.81,y,.51,.51,str(i+1),17,True,WHITE,PP_ALIGN.CENTER)
    txt(s,1.54,y-.03,2.40,.52,h,18,True)
    txt(s,4.03,y-.03,8.11,.57,b,17)
    if i<3: line(s,.95,y+.59,.95,y+1.08,PALE,2.1)
eq(s,'robust',1.53,6.54,5.64,.46)
txt(s,7.61,6.51,4.54,.46,'Σ 与半径仍需校准和论证。',16.4)
txt(s,8.32,5.95,3.68,.35,'相关方法 [6,7]',15.8)
note(s,'回答：202610-系统方案v2.pdf提出了什么，当前到哪一步。椭球、功率与公平效用、组稀疏、事件触发是拟研究组合，未接入当前集中式统计误差WMMSE–ADMM。当前条件协方差C不可直接称已标定椭球；最坏情况子问题、组稀疏优化、本地TRP共识与事件触发通信未实现。椭球参考[6] Zheng et al. 2008 DOI:10.1155/2008/609028；组稀疏参考[7] Shi et al., Group Sparse Beamforming for Green Cloud-RAN, IEEE TWC 2014, DOI:10.1109/TWC.2014.040214.131770。关闭部分TRP–UE链路不等于整TRP休眠。下一页说明机制成立还要解决的数学和系统问题。')

s=keep_title(prs.slides[13])
heading(s,.57,1.32,5.82,'代表目标与三项难点')
eq(s,'proposal',.79,1.96,11.52,.62)
for i,(h,b) in enumerate([('鲁棒内层','先定义误差集合与对偶/内层问题；联合 WMMSE 仍非凸。'),
                          ('稀疏更新','耦合二次项加组稀疏，不能宣称整体一次软阈值闭式求解。'),
                          ('本地共识','明确 TRP 局部变量、CPU 共识、消息缓存、触发与停止。')]):
    y=3.02+i*1.11
    rect(s,.76,y,2.17,.67,PALE,True)
    txt(s,.82,y+.04,2.05,.58,h,18,True,BLUE,PP_ALIGN.CENTER)
    txt(s,3.30,y+.03,8.76,.62,b,17.2)
    if i<2: line(s,3.30,y+.79,12.15,y+.79,PALE,1.0)
txt(s,.81,6.56,11.44,.47,'更大误差、功率下降和通信节省均需在同条件对照中验证。',17.1,True)
note(s,'回答：新版方案要从想法走到可运行需要哪些论证。公式是目标结构示意，精确效用U、罚项系数和误差集合尚待定义。WMMSE联合问题通常非凸；固定部分变量后的子问题凸性需逐项证明。一般耦合二次项加组稀疏不能声称整个子问题一次软阈值完成。TRP本地更新与事件触发还需定义通信缓存、共识更新时序及收敛条件。更大CSI误差不必然形成更宽波束；关闭一条服务链路不等于TRP休眠。功率或消息节省比例是待实测目标。依据202610-系统方案v2.pdf以及当前代码未实现边界。[7] Shi et al. 2014用于组稀疏方法背景，不证明当前代码效果。')

assert len(prs.slides) == 14
prs.save(OUT)
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == actual, 'Input was modified'
print('saved',OUT)
